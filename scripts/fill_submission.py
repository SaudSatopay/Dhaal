"""Fill the organizers' PPT template with Dhaal content -> pitch/Dhaal-HackX-submission.pptx
v2: bold typography + drawn diagrams (architecture flow, flywheel, stat tiles).

Re-run after editing TEAM_NAME / TEAM_ID:  python scripts/fill_submission.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEAM_NAME = "[Your Team Name]"   # <-- FILL from registration
TEAM_ID = "TEAM130"

INK = RGBColor(0x0E, 0x28, 0x41)
ACC = RGBColor(0x15, 0x60, 0x82)
RED = RGBColor(0xC6, 0x28, 0x28)
SAF = RGBColor(0xE8, 0xA1, 0x3B)
WHT = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0x51, 0x5E, 0x6E)

p = Presentation("pitch/template/PPT Template - MUJ HACKX 4.0.pptx")
S = list(p.slides)

def txbox(slide, x, y, w, h):
    b = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf

def para(tf, first):
    return tf.paragraphs[0] if first else tf.add_paragraph()

def run(pr, text, size, bold=False, color=INK, font="Calibri", align=None):
    r = pr.add_run(); r.text = text
    f = r.font; f.name = font; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    if align: pr.alignment = align
    return r

def lines(slide, x, y, w, h, blocks):
    tf = txbox(slide, x, y, w, h)
    first = True
    for kind, text in blocks:
        pr = para(tf, first); first = False
        if kind == "h":
            run(pr, text, 18, True, ACC, "Arial")
            pr.space_before = Pt(12); pr.space_after = Pt(3)
        elif kind == "big":
            run(pr, text, 23, True, INK, "Arial"); pr.space_after = Pt(6)
        elif kind == "bigred":
            run(pr, text, 25, True, RED, "Arial"); pr.space_after = Pt(6)
        elif kind == "red":
            run(pr, text, 16, True, RED, "Arial"); pr.space_before = Pt(10)
        else:
            run(pr, text, 14, False, INK); pr.space_after = Pt(4)
    return tf

def node(slide, x, y, w, h, fill, line_c, title, sub, tcolor, scolor, tsize=16, ssize=10.5):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.08
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line_c; sh.line.width = Pt(2.25)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12); tf.margin_top = tf.margin_bottom = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pr = tf.paragraphs[0]; pr.alignment = PP_ALIGN.CENTER
    run(pr, title, tsize, True, tcolor, "Arial")
    if sub:
        pr2 = tf.add_paragraph(); pr2.alignment = PP_ALIGN.CENTER
        run(pr2, sub, ssize, False, scolor)
    return sh

def arrow(slide, x, y, w=0.45, h=0.5, color=SAF):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False
    return a

def stat_tile(slide, x, y, w, h, num, label, ncolor=INK):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = WHT
    sh.line.color.rgb = INK; sh.line.width = Pt(2.25); sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08); tf.margin_top = Inches(0.08); tf.margin_bottom = Inches(0.06)
    pr = tf.paragraphs[0]; pr.alignment = PP_ALIGN.CENTER
    run(pr, num, 38, True, ncolor, "Arial")
    pr2 = tf.add_paragraph(); pr2.alignment = PP_ALIGN.CENTER
    run(pr2, label, 10.5, True, SOFT, "Arial")

# ================= Slide 1 — cover =================
sub = [sh for sh in S[0].shapes if sh.has_text_frame and "Team Name" in sh.text_frame.text][0]
tf = sub.text_frame
cover = [
    ("Dhaal (ढाल) — the pre-payment scam shield for every Indian", 20, True, ACC),
    (f"Team Name: {TEAM_NAME}", 16, False, INK),
    (f"Team ID: {TEAM_ID}", 16, False, INK),
    ("Problem Statement Selected: Consumer Protection Against Payment Scams (Fintech PS #7)", 16, False, INK),
    ("Theme: Fintech", 16, False, INK),
]
while len(tf.paragraphs) < len(cover):
    tf.add_paragraph()
for pr, (txt, size, bold, color) in zip(tf.paragraphs, cover):
    for r in list(pr.runs)[1:]:
        r._r.getparent().remove(r._r)
    r = pr.runs[0] if pr.runs else pr.add_run()
    r.text = txt
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Arial" if bold else "Calibri"

# ================= Slide 2 — IDEA (2x2 surface tiles) =================
s = S[1]
tfx = txbox(s, 0.9, 1.95, 11.5, 0.55)
pr = tfx.paragraphs[0]
run(pr, "पैसे भेजने से पहले — एक जाँच।  ", 24, True, INK, "Arial")
run(pr, "One check before you pay: paste it, photo it, or SPEAK it.", 17, True, ACC, "Arial")
tiles = [
    ("जाँच · Check", "Verdict + every weighted signal, in the user's language (हिं/EN toggle) — and spoken aloud. Never \"safe\", only \"no known risk\"."),
    ("परिवार की ढाल · Guardian", "An elderly user's risky payment pings family for a 10-second Allow/Block. Every check visible; dignity intact — risk, never transactions."),
    ("धोखों का नक्शा · Intel", "Report → human verify → everyone's live blocklist. A scam burned in Jaipur today can't work in Jodhpur tomorrow."),
    ("पहला घंटा · Recover", "1930 call script + cybercrime.gov.in complaint + bank letter — auto-drafted in Hindi, in the golden hour that gets money back."),
]
pos = [(0.9, 2.65), (6.85, 2.65), (0.9, 4.6), (6.85, 4.6)]
for (title, desc), (x, y) in zip(tiles, pos):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.55), Inches(1.8))
    sh.adjustments[0] = 0.055
    sh.fill.solid(); sh.fill.fore_color.rgb = WHT
    sh.line.color.rgb = INK; sh.line.width = Pt(2.5); sh.shadow.inherit = False
    tfi = sh.text_frame; tfi.word_wrap = True
    tfi.margin_left = tfi.margin_right = Inches(0.16); tfi.margin_top = Inches(0.1)
    pr = tfi.paragraphs[0]; run(pr, title, 17, True, ACC, "Arial"); pr.space_after = Pt(4)
    pr2 = tfi.add_paragraph(); run(pr2, desc, 12, False, INK)
tfz = txbox(s, 0.9, 6.62, 11.5, 0.5)
pr = tfz.paragraphs[0]
run(pr, "The verdict is computed by rules — the AI explains it, and can never change it.   ", 15, True, INK, "Arial")
run(pr, "LIVE: dhaal-delta.vercel.app — check your own inbox in Q&A.", 15, True, RED, "Arial")

# ================= Slide 3 — TECHNICAL (architecture diagram) =================
s = S[2]
lines(s, 0.9, 1.95, 11.5, 0.75, [
    ("t", "Next.js PWA (Vercel) · FastAPI on Vercel Python · MongoDB Atlas · Claude claude-sonnet-5 · Sarvam AI (Saarika ASR + Bulbul TTS) · jsQR on-device"),
])
node(s, 0.9, 2.75, 2.5, 1.55, WHT, INK, "INPUT", "text · URL · UPI · QR (on-device) · voice ASR", INK, SOFT, 15, 10)
arrow(s, 3.5, 3.28)
node(s, 4.05, 2.6, 4.5, 1.85, RED, RED, "SIGNAL ENGINE", "deterministic verdict — collect-vs-pay · lookalike domains · 10 scam-script families · coercion · bait-VPA · blocklist  |  ≥60 खतरा · ≥30 सावधान", WHT, RGBColor(0xF3, 0xD9, 0xD9), 17, 10)
arrow(s, 8.65, 3.28)
node(s, 9.2, 2.75, 3.2, 1.55, WHT, INK, "EXPLAIN & SPEAK", "Claude narrates from signals — zero verdict weight · Bulbul speaks the warning", INK, SOFT, 14, 10)
conn = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.27), Inches(4.45), Inches(0.06), Inches(0.5))
conn.fill.solid(); conn.fill.fore_color.rgb = INK; conn.line.fill.background(); conn.shadow.inherit = False
node(s, 4.85, 4.95, 2.9, 0.8, WHT, ACC, "MongoDB Atlas", "durable community intel — human-verified blocklist", ACC, SOFT, 13, 9.5)
lines(s, 0.9, 6.0, 11.5, 1.1, [
    ("h", "Reliability engineering"),
    ("t", "Retry-once + deterministic fallback on every external call (dead Wi-Fi degrades prose, never protection; mocked responses say mocked:true) · 49 automated checks green · measured live: verdict in ms, p50 5.4s with explanation, Hindi ASR 0.6s."),
])

# ================= Slide 4 — UX =================
s = S[3]
lines(s, 0.9, 2.0, 11.5, 4.4, [
    ("h", "A public-safety poster, not a SaaS dashboard"),
    ("t", "Hand-built \"suraksha poster\" identity — ink + saffron hazard language, Anek Devanagari display type, a rubber-stamp community count on the verdict card."),
    ("h", "Built for the next billion users"),
    ("t", "Hindi-first with a global हिं/EN toggle that holds across every page and every output — including the spoken warning · voice in AND out for low-literacy users · QR photos decoded on-device (the image never leaves the phone) · entry-level phones, weak networks."),
    ("h", "Dignity by design"),
    ("t", "Guardian mode shares the risk summary only — never transactions. Risky checks ask for a decision; clean ones appear as quiet notes. Honest language everywhere."),
])
tfq = txbox(s, 0.9, 6.5, 11.5, 0.6)
pr = tfq.paragraphs[0]
run(pr, "Danger screams. Safety whispers.", 24, True, RED, "Arial")

# ================= Slide 5 — FEASIBILITY (stat tiles) =================
s = S[4]
tfb = txbox(s, 0.9, 1.95, 11.5, 0.55)
run(tfb.paragraphs[0], "Not a concept — deployed, tested and durable tonight.", 23, True, INK, "Arial")
stat_tile(s, 0.9, 2.7, 2.72, 1.5, "49", "AUTOMATED CHECKS GREEN")
stat_tile(s, 3.83, 2.7, 2.72, 1.5, "5.4s", "P50 CHECK, LIVE PROD", ACC)
stat_tile(s, 6.76, 2.7, 2.72, 1.5, "450+", "VERIFIED REPORTS IN ATLAS", ACC)
stat_tile(s, 9.69, 2.7, 2.72, 1.5, "₹0", "INFRA — FREE TIERS", RED)
lines(s, 0.9, 4.55, 11.5, 2.5, [
    ("h", "Live now"),
    ("t", "dhaal-delta.vercel.app (app) · dhaal-api.vercel.app (API) · guardian pairing, community flywheel and the full voice loop verified end-to-end on stage hardware."),
    ("h", "Risks & mitigations"),
    ("t", "API outage → deterministic fallbacks, demo survives · false positives → delivery-vs-request context rules, tested on real bank OTP/debit SMS (score 0) · blocklist poisoning → human verification before any indicator goes live · scale → stateless serverless + Atlas, TTS caching."),
])

# ================= Slide 6 — IMPACT (flywheel diagram) =================
s = S[5]
tfr = txbox(s, 0.9, 1.95, 11.5, 0.9)
pr = tfr.paragraphs[0]
run(pr, "₹22,495 crore", 30, True, RED, "Arial")
run(pr, "  lost to cyber fraud in 2025 — and the victim pressed PAY themselves.", 19, True, INK, "Arial")
lines(s, 0.9, 2.85, 11.5, 0.6, [
    ("t", "24.51B UPI txns/month (NPCI Aug '26) · 28.1 lakh complaints in 2025 (I4C) · ₹805 crore UPI fraud in FY26 (Parliament) — bank controls cannot stop payments the victim authorises."),
])
fw = [("1 REPORT", "10 seconds"), ("HUMAN VERIFY", "no rumour poisoning"), ("BLOCKLIST LIVE", "durable · Atlas"), ("EVERYONE SHIELDED", "every phone, instantly")]
x = 0.9
for i, (t1, t2) in enumerate(fw):
    node(s, x, 3.85, 2.55, 1.05, WHT if i < 3 else RGBColor(0xE8, 0xF3, 0xEA), INK if i < 3 else RGBColor(0x2E, 0x7D, 0x32), t1, t2, INK if i < 3 else RGBColor(0x2E, 0x7D, 0x32), SOFT, 13.5, 9.5)
    if i < 3:
        arrow(s, x + 2.6, 4.18, 0.38, 0.42)
    x += 3.0
tfl = txbox(s, 0.9, 5.1, 11.5, 0.4)
run(tfl.paragraphs[0], "↻ herd immunity for fraud — the shield compounds with every user", 14, True, ACC, "Arial")
lines(s, 0.9, 5.7, 11.5, 1.4, [
    ("h", "Who benefits"),
    ("t", "Every Indian with a phone (free forever) · the elderly — the most-targeted group (guardian mode) · banks & UPI apps (fewer reimbursements) · cyber cells (structured early reports — fast 1930 reporting has already saved ₹7,130+ crore)."),
])

# ================= Slide 7 — BUSINESS =================
s = S[6]
lines(s, 0.9, 2.0, 11.5, 4.3, [
    ("big", "Citizens: free forever. Institutions pay."),
    ("t", "Protection this fundamental must not be paywalled — and by construction the marginal cost per check is a fraction of a paisa (deterministic verdict; one short narration)."),
    ("h", "Revenue — from those who bear today's fraud cost"),
    ("t", "Risk-API + verified-blocklist licensing to banks, fintechs and UPI apps (they reimburse this fraud today; prevention is cheaper) · pre-payment check SDK inside UPI apps — the check runs before the PAY button renders · deployment partnerships: state cyber cells & the CSC network."),
    ("h", "Sustainability & moat"),
    ("t", "The community blocklist compounds with every user — a network effect no new entrant starts with · Roadmap: one-tap 1930/I4C hand-off · 10+ Indic languages via the same Sarvam pipeline (one parameter) · auto-verify thresholds for moderation at scale."),
])

# ================= Slide 8 — REFERENCES =================
s = S[7]
lines(s, 0.9, 2.0, 11.5, 5.0, [
    ("h", "Data sources (verified 11 Sep 2026)"),
    ("t", "NPCI monthly UPI statistics — 24.51B transactions / ₹29.82 lakh crore, Aug 2026 · I4C (MHA): ₹22,495 crore lost, 28.1 lakh complaints, 2025; 32.4M calls to 1930 · Ministry of Finance, Rajya Sabha reply: ₹805 crore UPI fraud, 10.64 lakh incidents, FY26 (to Nov) · PIB/MHA: ₹7,130+ crore saved via CFCFRMS (1930) early reporting · RBI guidelines on unauthorised-transaction liability."),
    ("h", "Domain research"),
    ("t", "Scam-script corpus from I4C advisories & documented fraud patterns: digital arrest, KYC-expiry, collect-request refund bait, OLX/army advance, job-task & loan-fee scams, electricity disconnection, victim-voiced coercion · upi:// deep-link collect-vs-pay semantics · NPCI PSP handle registry for VPA verification."),
    ("h", "Build"),
    ("t", "Live app: dhaal-delta.vercel.app · API: dhaal-api.vercel.app · Repo: github.com/SaudSatopay/MUJ-HACKX (49 automated checks, full contracts & docs) · Sponsor tech in production: Sarvam AI (Saarika ASR + Bulbul TTS), MongoDB Atlas."),
])

out = "pitch/Dhaal-HackX-submission.pptx"
try:
    p.save(out)
except PermissionError:
    out = "pitch/Dhaal-HackX-submission-v2.pptx"
    p.save(out)
print("saved", out)
